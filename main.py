from __future__ import annotations

import csv
import io
import json
import os
import re
import sys
import threading
import traceback
import webbrowser
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import barcode
import qrcode
import tkinter as tk
from barcode.writer import ImageWriter
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageTk
from tkinter import filedialog, messagebox, ttk

from feishu_client import FeishuError, begin_user_auth, complete_user_auth, environment_status, load_feishu_data

try:
    import win32print
except ImportError:
    win32print = None


APP_NAME = "飞书 Excel 标签打印助手"
APP_VERSION = "0.4.2"
DP_MM = 12  # 300 dpi is approximately 12 dots/mm

CONFIG_DIR = Path(os.environ.get("APPDATA", Path.home())) / "PrimeHandLabelBridge"
CONFIG_FILE = CONFIG_DIR / "config.json"

TEMPLATES = {
    "零件标签 40×30 mm": {"kind": "simple", "width_mm": 40, "height_mm": 30},
}

FIELDS = [
    ("code", "物料编码/PN"),
    ("name", "物料名称"),
    ("supplier", "供应商"),
    ("rev", "版本/REV"),
    ("lot", "批次/LOT"),
    ("sn", "序列号/SN"),
    ("link", "飞书记录链接"),
    ("qty", "打印数量"),
    ("firmware", "固件版本"),
]

ALIASES = {
    "code": ["料号", "物料编码", "物料编号", "零件号", "pn", "partnumber", "编码", "code"],
    "name": ["零件", "物料名称", "零件名称", "名称", "品名", "name"],
    "supplier": ["供应商", "厂家", "厂商", "supplier", "vendor"],
    "rev": ["版本", "物料版本", "设计版本", "rev", "revision"],
    "lot": ["批次", "批次号", "lot", "lotnumber"],
    "sn": ["sn", "序列号", "整机sn", "产品sn", "serialnumber"],
    "link": ["记录分享链接", "分享链接", "记录链接", "飞书链接", "链接", "recordlink", "url", "二维码内容", "二维码"],
    "qty": ["打印数量", "标签数量", "数量", "qty", "copies", "打印份数"],
    "firmware": ["固件版本", "软件版本", "firmware", "fw"],
}


def resource_path(relative: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / relative


def load_config() -> dict[str, Any]:
    try:
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_config(config: dict[str, Any]) -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    CONFIG_FILE.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return " / ".join(part for part in (clean(item) for item in value) if part)
    if isinstance(value, dict):
        for key in ("text", "name", "value"):
            if key in value:
                return clean(value[key])
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized(value: Any) -> str:
    return re.sub(r"[\s_\-/（）()]+", "", clean(value)).lower()


def auto_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {header: normalized(header) for header in headers}
    result: dict[str, str] = {}
    for field, aliases in ALIASES.items():
        exact = [h for h, nh in normalized_headers.items() if nh in aliases]
        fuzzy = [h for h, nh in normalized_headers.items() if any(a in nh for a in aliases)]
        result[field] = (exact or fuzzy or [""])[0]
    return result


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = []
    if bold:
        candidates.extend([
            Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/msyhbd.ttc",
            Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/simhei.ttf",
        ])
    candidates.extend([
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/msyh.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/simsun.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/arial.ttf",
    ])
    for font_path in candidates:
        if font_path.exists():
            try:
                return ImageFont.truetype(str(font_path), size=size)
            except Exception:
                continue
    return ImageFont.load_default()


def fit_text(draw: ImageDraw.ImageDraw, text: str, max_width: int, start_size: int, min_size: int = 16, bold: bool = False):
    text = clean(text)
    for size in range(start_size, min_size - 1, -1):
        font = get_font(size, bold)
        if draw.textbbox((0, 0), text, font=font)[2] <= max_width:
            return font
    return get_font(min_size, bold)


def make_qr(payload: str, max_size: int) -> Image.Image:
    qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=8, border=2)
    qr.add_data(payload)
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white").convert("1")
    image.thumbnail((max_size, max_size), Image.Resampling.NEAREST)
    return image


def make_code128(payload: str, width: int, height: int) -> Image.Image | None:
    payload = clean(payload)
    if not payload:
        return None
    out = io.BytesIO()
    code = barcode.get("code128", payload, writer=ImageWriter())
    code.write(out, options={
        "module_width": 0.20,
        "module_height": 7.0,
        "quiet_zone": 0.8,
        "font_size": 0,
        "text_distance": 0,
        "write_text": False,
        "dpi": 300,
    })
    out.seek(0)
    image = Image.open(out).convert("1")
    image.thumbnail((width, height), Image.Resampling.LANCZOS)
    return image


@dataclass
class LabelRecord:
    source_row: int
    code: str = ""
    name: str = ""
    supplier: str = ""
    rev: str = ""
    lot: str = ""
    sn: str = ""
    link: str = ""
    qty: str = ""
    firmware: str = ""
    raw_fields: dict[str, str] = field(default_factory=dict)

    def primary(self, kind: str) -> str:
        if kind == "product":
            return self.sn or self.code
        return self.code or self.sn

    def qr_payload(self, kind: str) -> str:
        return self.link or self.primary(kind)


def centered_x(draw: ImageDraw.ImageDraw, text: str, font, canvas_width: int) -> int:
    bbox = draw.textbbox((0, 0), text, font=font)
    return max(0, (canvas_width - (bbox[2] - bbox[0])) // 2)


def render_label(
    record: LabelRecord,
    template_name: str,
    display_columns: list[tuple[str, str]] | None = None,
    inbound_date: str = "",
) -> Image.Image:
    template = TEMPLATES[template_name]
    kind = template["kind"]
    width = int(template["width_mm"] * DP_MM)
    height = int(template["height_mm"] * DP_MM)
    image = Image.new("1", (width, height), 1)
    draw = ImageDraw.Draw(image)
    margin = max(10, int(width * 0.025))
    draw.rectangle((1, 1, width - 2, height - 2), outline=0, width=2)

    if kind == "simple":
        # Keep all content inside a 2.5 mm safe area. This avoids the weak left edge
        # of a narrow ribbon and centers the full layout on the physical label.
        safe_margin = int(2.5 * DP_MM)
        content_width = width - safe_margin * 2
        lines = display_columns or [
            ("零件", record.name),
            ("料号", record.code),
            ("供应商", record.supplier),
        ]
        lines = lines[:5]
        if inbound_date:
            lines.append(("入库日期", inbound_date))
        compact = len(lines) > 3
        y = safe_margin
        for index, (header, value) in enumerate(lines):
            value = clean(value) or "—"
            text = value if index == 0 else f"{header}：{value}"
            start_size = (32 if index == 0 else 20) if compact else (40 if index == 0 else 28)
            min_size = 14 if compact else 16
            font = fit_text(draw, text, content_width, start_size, min_size, index == 0)
            draw.text((centered_x(draw, text, font, width), y), text, font=font, fill=0)
            y += draw.textbbox((0, 0), text, font=font)[3] + (4 if index == 0 else 1)

        qr = make_qr(record.link, 108) if record.link else None
        gap = 8 if qr else 0
        barcode_max_width = content_width - ((qr.width + gap) if qr else 0)
        bar = make_code128(record.code, barcode_max_width, 88) if record.code else None
        group_width = (bar.width if bar else 0) + (gap if bar and qr else 0) + (qr.width if qr else 0)
        graphic_height = max(bar.height if bar else 0, qr.height if qr else 0)
        gx = max(safe_margin, (width - group_width) // 2)
        bottom_limit = height - safe_margin - graphic_height
        gy = min(bottom_limit, max(y + 10, int(height * 0.48)))
        if bar is not None:
            image.paste(bar, (gx, gy + (graphic_height - bar.height) // 2))
            gx += bar.width + (gap if qr else 0)
        if qr is not None:
            image.paste(qr, (gx, gy + (graphic_height - qr.height) // 2))
        return image

    qr_size = min(height - margin * 2, int(width * (0.34 if kind != "product" else 0.38)))
    qr_payload = record.qr_payload(kind)
    if qr_payload:
        qr = make_qr(qr_payload, qr_size)
        qx = width - margin - qr.width
        qy = margin
        image.paste(qr, (qx, qy))
    else:
        qx = width - margin - qr_size
        draw.rectangle((qx, margin, width - margin, margin + qr_size), outline=0, width=2)
        draw.line((qx, margin, width - margin, margin + qr_size), fill=0, width=2)
        draw.line((width - margin, margin, qx, margin + qr_size), fill=0, width=2)

    left_width = qx - margin * 2
    primary = record.primary(kind) or "未映射编码"
    primary_font = fit_text(draw, primary, left_width, 38 if kind == "product" else 34, 18, True)
    draw.text((margin, margin), primary, font=primary_font, fill=0)
    y = margin + draw.textbbox((0, 0), primary, font=primary_font)[3] + 3

    name_font = fit_text(draw, record.name, left_width, 24, 15, False)
    if record.name:
        draw.text((margin, y), record.name, font=name_font, fill=0)
        y += draw.textbbox((0, 0), record.name, font=name_font)[3] + 3

    details = []
    if record.rev:
        details.append(f"REV {record.rev}")
    if record.lot:
        details.append(f"LOT {record.lot}")
    detail_lines = ["  |  ".join(details)] if details else []
    secondary = []
    if kind != "product" and record.sn:
        secondary.append(f"SN {record.sn}")
    if record.firmware:
        secondary.append(f"FW {record.firmware}")
    if secondary:
        detail_lines.append("  |  ".join(secondary))
    for detail_text in detail_lines:
        detail_font = fit_text(draw, detail_text, left_width, 18, 10, False)
        draw.text((margin, y), detail_text, font=detail_font, fill=0)
        y += draw.textbbox((0, 0), detail_text, font=detail_font)[3] + 1

    barcode_height = max(55, int(height * 0.24))
    barcode_width = left_width
    bar = make_code128(primary, barcode_width, barcode_height)
    if bar is not None:
        bx = margin + max(0, (barcode_width - bar.width) // 2)
        by = height - margin - bar.height
        image.paste(bar, (bx, by))
    return image


def image_to_zpl(image: Image.Image) -> bytes:
    mono = image.convert("1")
    width, height = mono.size
    row_bytes = (width + 7) // 8
    data = bytearray()
    pixels = mono.load()
    for y in range(height):
        for byte_x in range(row_bytes):
            value = 0
            for bit in range(8):
                x = byte_x * 8 + bit
                is_black = x < width and pixels[x, y] == 0
                if is_black:
                    value |= 1 << (7 - bit)
            data.append(value)
    hex_data = data.hex().upper()
    zpl = f"^XA^PW{width}^LL{height}^LH0,0^FO0,0^GFA,{len(data)},{len(data)},{row_bytes},{hex_data}^FS^XZ\r\n"
    return zpl.encode("ascii")


def list_printers() -> list[str]:
    if win32print is None:
        return []
    flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
    return sorted({p[2] for p in win32print.EnumPrinters(flags)}, key=str.lower)


def send_raw(printer_name: str, jobs: list[bytes], document_name: str = "Prime Hand Labels") -> None:
    if win32print is None:
        raise RuntimeError("缺少 Windows 打印组件 pywin32。")
    handle = win32print.OpenPrinter(printer_name)
    try:
        job = win32print.StartDocPrinter(handle, 1, (document_name, None, "RAW"))
        try:
            win32print.StartPagePrinter(handle)
            for payload in jobs:
                win32print.WritePrinter(handle, payload)
            win32print.EndPagePrinter(handle)
        finally:
            win32print.EndDocPrinter(handle)
    finally:
        win32print.ClosePrinter(handle)


class LabelBridgeApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_NAME}  v{APP_VERSION}")
        self.geometry("1220x800")
        self.minsize(1050, 700)
        self.option_add("*Font", ("Microsoft YaHei UI", 10))

        self.config_data = load_config()
        self.file_path = tk.StringVar(value="")
        self.sheet_name = tk.StringVar(value="")
        self.header_row = tk.IntVar(value=1)
        saved_template = self.config_data.get("template")
        self.template_name = tk.StringVar(value=saved_template if saved_template in TEMPLATES else next(iter(TEMPLATES)))
        self.printer_name = tk.StringVar(value=self.config_data.get("printer", ""))
        self.fixed_copies = tk.IntVar(value=1)
        self.use_excel_qty = tk.BooleanVar(value=False)
        self.status_text = tk.StringVar(value="请选择 Excel 或我为你导出的飞书打印数据包。")
        self.mapping_vars = {field: tk.StringVar(value="") for field, _ in FIELDS}
        self.display_column_vars = [tk.StringVar(value="") for _ in range(5)]
        self.feishu_url = tk.StringVar(value="")
        self.inbound_date = tk.StringVar(value=date.today().isoformat())

        self.raw_rows: list[dict[str, Any]] = []
        self.headers: list[str] = []
        self.records: list[LabelRecord] = []
        self.table_headers: list[str] = ["零件", "料号", "供应商"]
        self.preview_photo: ImageTk.PhotoImage | None = None

        self._build_ui()
        self.refresh_printers()

    def _build_ui(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(2, weight=1)

        file_box = ttk.LabelFrame(self, text="1. 选择 Excel / 飞书打印数据包")
        file_box.grid(row=0, column=0, padx=12, pady=(10, 5), sticky="ew")
        file_box.columnconfigure(1, weight=1)
        ttk.Button(file_box, text="选择文件", command=self.choose_file).grid(row=0, column=0, padx=8, pady=8)
        ttk.Entry(file_box, textvariable=self.file_path, state="readonly").grid(row=0, column=1, padx=4, pady=8, sticky="ew")
        ttk.Label(file_box, text="Excel 读取第一个工作表；飞书数据包自带逐行分享链接").grid(row=0, column=2, padx=12)
        ttk.Label(file_box, text="飞书链接").grid(row=1, column=0, padx=8, pady=(2, 8))
        ttk.Entry(file_box, textvariable=self.feishu_url).grid(row=1, column=1, padx=4, pady=(2, 8), sticky="ew")
        feishu_actions = ttk.Frame(file_box)
        feishu_actions.grid(row=1, column=2, padx=8, pady=(2, 8))
        self.feishu_load_button = ttk.Button(feishu_actions, text="读取飞书", command=self.load_feishu_link)
        self.feishu_load_button.pack(side="left", padx=3)
        ttk.Button(feishu_actions, text="扫码授权", command=self.start_feishu_auth).pack(side="left", padx=3)
        ttk.Button(feishu_actions, text="检查环境", command=self.check_feishu_environment).pack(side="left", padx=3)

        controls = ttk.LabelFrame(self, text="2. 选择打印机和数量")
        controls.grid(row=1, column=0, padx=12, pady=5, sticky="ew")
        ttk.Label(controls, text="打印机").grid(row=0, column=0, padx=(10, 3), pady=8)
        self.printer_combo = ttk.Combobox(controls, textvariable=self.printer_name, state="readonly", width=34)
        self.printer_combo.grid(row=0, column=1, padx=4)
        ttk.Button(controls, text="刷新打印机", command=self.refresh_printers).grid(row=0, column=2, padx=8)
        ttk.Label(controls, text="每种打印份数").grid(row=0, column=3, padx=(20, 3))
        ttk.Spinbox(controls, from_=1, to=999, textvariable=self.fixed_copies, width=6).grid(row=0, column=4, padx=4)
        ttk.Label(controls, text="标签：40×30 mm / 300 dpi").grid(row=0, column=5, padx=20)
        ttk.Label(controls, text="标签显示列").grid(row=1, column=0, padx=(10, 3), pady=(2, 4))
        for index, var in enumerate(self.display_column_vars):
            combo = ttk.Combobox(controls, textvariable=var, state="readonly", width=18)
            row = 1 if index < 3 else 2
            column = index + 1 if index < 3 else index - 2
            combo.grid(row=row, column=column, padx=4, pady=(2, 4 if row == 1 else 8))
            combo.bind("<<ComboboxSelected>>", lambda _e: self.show_preview())
            setattr(self, f"display_combo_{index}", combo)
        ttk.Label(controls, text="最多选择 5 列；分享链接只生成二维码").grid(row=1, column=4, columnspan=2, padx=15)
        ttk.Label(controls, text="入库日期").grid(row=2, column=3, padx=(15, 3), pady=(2, 8))
        inbound_entry = ttk.Entry(controls, textvariable=self.inbound_date, width=13)
        inbound_entry.grid(row=2, column=4, padx=4, pady=(2, 8))
        inbound_entry.bind("<Return>", lambda _e: self.commit_inbound_date())
        inbound_entry.bind("<FocusOut>", lambda _e: self.commit_inbound_date())
        ttk.Button(controls, text="今天", command=self.set_inbound_today).grid(row=2, column=5, padx=4, pady=(2, 8))

        main = ttk.Panedwindow(self, orient="horizontal")
        main.grid(row=2, column=0, padx=12, pady=5, sticky="nsew")
        table_frame = ttk.Frame(main)
        preview_frame = ttk.LabelFrame(main, text="标签预览（单击左侧记录切换）")
        main.add(table_frame, weight=4)
        main.add(preview_frame, weight=2)

        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)
        columns = ("row", "零件", "料号", "供应商")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="extended")
        headings = {"row": "Excel行", "零件": "零件", "料号": "料号", "供应商": "供应商"}
        widths = {"row": 70, "零件": 300, "料号": 180, "供应商": 180}
        for key in columns:
            self.tree.heading(key, text=headings[key])
            self.tree.column(key, width=widths[key], minwidth=50, anchor="w")
        scroll_y = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        scroll_x = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        scroll_y.grid(row=0, column=1, sticky="ns")
        scroll_x.grid(row=1, column=0, sticky="ew")
        self.tree.bind("<<TreeviewSelect>>", lambda _e: self.show_preview())

        preview_frame.rowconfigure(0, weight=1)
        preview_frame.columnconfigure(0, weight=1)
        self.preview_canvas = tk.Canvas(preview_frame, bg="#e8e8e8", highlightthickness=0)
        self.preview_canvas.grid(row=0, column=0, padx=8, pady=8, sticky="nsew")
        self.preview_canvas.bind("<Configure>", lambda _e: self.show_preview())

        action = ttk.Frame(self)
        action.grid(row=3, column=0, padx=12, pady=(4, 8), sticky="ew")
        ttk.Button(action, text="导出测试文件", command=lambda: self.export_zpl(False)).pack(side="left", padx=4)
        ttk.Button(action, text="全选行", command=self.select_all_rows).pack(side="left", padx=4)
        ttk.Button(action, text="清空选择", command=self.clear_row_selection).pack(side="left", padx=4)
        ttk.Button(action, text="打印选中的零件", command=lambda: self.print_records(False)).pack(side="right", padx=4)
        ttk.Button(action, text="打印全部零件", command=lambda: self.print_records(True)).pack(side="right", padx=4)
        ttk.Label(action, textvariable=self.status_text).pack(side="left", padx=15)

    def choose_file(self):
        path = filedialog.askopenfilename(
            title="选择飞书导出的表格",
            filetypes=[("支持的文件", "*.xlsx *.xlsm *.csv *.json"), ("Excel 文件", "*.xlsx *.xlsm"), ("飞书打印数据包", "*.json"), ("CSV 文件", "*.csv"), ("所有文件", "*.*")],
        )
        if not path:
            return
        self.file_path.set(path)
        try:
            if Path(path).suffix.lower() == ".json":
                self.sheet_name.set("飞书打印数据")
            elif Path(path).suffix.lower() == ".csv":
                self.sheet_name.set("CSV")
            else:
                workbook = load_workbook(path, read_only=True, data_only=True)
                self.sheet_name.set(workbook.sheetnames[0])
                workbook.close()
            self.load_sheet()
        except Exception as exc:
            self.show_error("无法读取文件", exc)

    def check_feishu_environment(self):
        self.status_text.set("正在检查飞书授权环境…")

        def worker():
            try:
                status = environment_status()
                user = status.get("identities", {}).get("user", {})
                name = user.get("userName") or "当前用户"
                state = user.get("status", "unknown")
                message = f"lark-cli 可用；{name}，用户授权状态：{state}"
                self.after(0, lambda: messagebox.showinfo(APP_NAME, message))
                self.after(0, lambda: self.status_text.set(message))
            except Exception as exc:
                self.after(0, lambda: self.show_error("飞书环境不可用", exc))

        threading.Thread(target=worker, daemon=True).start()

    def start_feishu_auth(self):
        self.status_text.set("正在发起飞书扫码授权…")

        def begin_worker():
            try:
                data = begin_user_auth()
                verification_url = data.get("verification_url") or data.get("verification_uri_complete")
                device_code = data.get("device_code")
                if not verification_url or not device_code:
                    raise FeishuError("飞书没有返回授权链接或 device_code。")
                self.after(0, lambda: self.show_auth_dialog(verification_url, device_code))
            except Exception as exc:
                self.after(0, lambda: self.show_error("无法发起飞书授权", exc))

        threading.Thread(target=begin_worker, daemon=True).start()

    def show_auth_dialog(self, verification_url: str, device_code: str):
        dialog = tk.Toplevel(self)
        dialog.title("飞书扫码授权")
        dialog.resizable(False, False)
        dialog.transient(self)
        qr_image = make_qr(verification_url, 280).convert("RGB")
        qr_photo = ImageTk.PhotoImage(qr_image)
        qr_label = ttk.Label(dialog, image=qr_photo)
        qr_label.image = qr_photo
        qr_label.pack(padx=24, pady=(20, 8))
        ttk.Label(dialog, text="请用飞书扫码，并确认 Base 权限授权").pack(padx=20, pady=4)
        auth_status = tk.StringVar(value="等待扫码授权…")
        ttk.Label(dialog, textvariable=auth_status).pack(padx=20, pady=4)
        ttk.Button(dialog, text="在浏览器打开", command=lambda: webbrowser.open(verification_url)).pack(pady=(4, 18))

        def complete_worker():
            try:
                complete_user_auth(device_code)
                self.after(0, lambda: auth_status.set("授权成功，可以关闭此窗口并读取飞书链接。"))
                self.after(0, lambda: self.status_text.set("飞书用户授权成功"))
            except Exception as exc:
                self.after(0, lambda: auth_status.set(f"授权未完成：{exc}"))

        threading.Thread(target=complete_worker, daemon=True).start()

    def load_feishu_link(self):
        url = self.feishu_url.get().strip()
        if not url:
            messagebox.showinfo(APP_NAME, "请先粘贴飞书多维表格的完整链接。")
            return
        self.feishu_load_button.configure(state="disabled")

        def progress(message: str):
            self.after(0, lambda m=message: self.status_text.set(m))

        def worker():
            try:
                data = load_feishu_data(url, progress)
                self.after(0, lambda: self.apply_feishu_data(data))
            except Exception as exc:
                self.after(0, lambda: self.show_error("读取飞书失败", exc))
            finally:
                self.after(0, lambda: self.feishu_load_button.configure(state="normal"))

        threading.Thread(target=worker, daemon=True).start()

    def apply_feishu_data(self, data: dict):
        headers = [clean(header) for header in data.get("fields", [])]
        headers.append("记录分享链接")
        self.raw_rows = []
        for item in data.get("records", []):
            values = {header: item.get("fields", {}).get(header, "") for header in headers}
            values["记录分享链接"] = item.get("record_share_link", "")
            self.raw_rows.append(values)
        self.file_path.set(f"飞书：{data.get('title', '多维表格')} / {data.get('table_id', '')}")
        self.sheet_name.set("飞书在线视图")
        self.finish_loaded_rows(headers)
        linked = sum(bool(record.link) for record in self.records)
        self.status_text.set(f"已直接读取飞书：{len(self.records)} 行，{linked} 个二维码链接")

    def _read_csv(self, path: str) -> list[list[Any]]:
        last_error = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                with open(path, "r", newline="", encoding=encoding) as handle:
                    return list(csv.reader(handle))
            except UnicodeDecodeError as exc:
                last_error = exc
        raise last_error or RuntimeError("无法识别 CSV 编码")

    def load_sheet(self):
        path = self.file_path.get()
        if not path:
            return
        try:
            header_index = max(0, int(self.header_row.get()) - 1)
            suffix = Path(path).suffix.lower()
            if suffix == ".json":
                package = json.loads(Path(path).read_text(encoding="utf-8"))
                headers = [clean(h) for h in package.get("fields", [])]
                if "记录分享链接" not in headers:
                    headers.append("记录分享链接")
                self.raw_rows = []
                for item in package.get("records", []):
                    values = {header: item.get("fields", {}).get(header, "") for header in headers}
                    values["记录分享链接"] = item.get("record_share_link", "")
                    self.raw_rows.append(values)
            elif suffix == ".csv":
                matrix = self._read_csv(path)
                if header_index >= len(matrix):
                    raise ValueError("标题行超出 CSV 数据范围。")
                headers = [clean(v) or f"未命名列{i + 1}" for i, v in enumerate(matrix[header_index])]
                rows = matrix[header_index + 1:]
                self.raw_rows = [dict(zip(headers, row + [""] * (len(headers) - len(row)))) for row in rows]
            else:
                workbook = load_workbook(path, read_only=False, data_only=True)
                sheet = workbook[self.sheet_name.get()]
                matrix = list(sheet.iter_rows())
                if header_index >= len(matrix):
                    workbook.close()
                    raise ValueError("标题行超出工作表数据范围。")
                headers = [clean(cell.value) or f"未命名列{i + 1}" for i, cell in enumerate(matrix[header_index])]
                self.raw_rows = []
                for row in matrix[header_index + 1:]:
                    values = {}
                    for i, header in enumerate(headers):
                        if i >= len(row):
                            values[header] = ""
                            continue
                        cell = row[i]
                        value = cell.value
                        if cell.hyperlink:
                            values[f"__hyperlink__{header}"] = cell.hyperlink.target
                        values[header] = value
                    self.raw_rows.append(values)
                workbook.close()
            self.finish_loaded_rows(headers)
            self.status_text.set(f"已读取 {self.sheet_name.get()}：{len(self.records)} 条可打印记录")
        except Exception as exc:
            self.show_error("读取工作表失败", exc)

    def finish_loaded_rows(self, headers: list[str]):
        self.headers = headers
        detected = auto_mapping(headers)
        for field, _label in FIELDS:
            self.mapping_vars[field].set(detected[field])
        link_header = self.mapping_vars["link"].get()
        selectable_headers = [h for h in headers if h != link_header]
        options = ["<不打印>"] + selectable_headers
        preferred = [self.mapping_vars["name"].get(), self.mapping_vars["code"].get(), self.mapping_vars["supplier"].get()]
        defaults = [value for value in preferred if value in selectable_headers]
        defaults.extend(header for header in selectable_headers if header not in defaults)
        for index, var in enumerate(self.display_column_vars):
            getattr(self, f"display_combo_{index}")["values"] = options
            var.set(defaults[index] if index < min(3, len(defaults)) else "<不打印>")
        self.apply_mapping()

    def select_all_rows(self):
        self.tree.selection_set(self.tree.get_children())
        self.show_preview()

    def clear_row_selection(self):
        self.tree.selection_remove(self.tree.selection())

    def set_inbound_today(self):
        self.inbound_date.set(date.today().isoformat())
        self.show_preview()

    def commit_inbound_date(self):
        if self.normalize_inbound_date(show_error=True):
            self.show_preview()

    def normalize_inbound_date(self, show_error: bool = True) -> str:
        raw = self.inbound_date.get().strip()
        if not raw:
            if show_error:
                messagebox.showwarning(APP_NAME, "请输入入库日期，格式为 YYYY-MM-DD。")
            return ""
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"):
            try:
                normalized_value = datetime.strptime(raw, pattern).date().isoformat()
                self.inbound_date.set(normalized_value)
                return normalized_value
            except ValueError:
                continue
        if show_error:
            messagebox.showwarning(APP_NAME, "入库日期格式不正确，请使用 YYYY-MM-DD，例如 2026-08-27。")
        return ""

    def apply_mapping(self):
        if not self.raw_rows:
            return
        self.records = []
        header_row = int(self.header_row.get())
        for offset, raw in enumerate(self.raw_rows, start=header_row + 1):
            values = {}
            for field, _label in FIELDS:
                header = self.mapping_vars[field].get()
                if field == "link" and header:
                    values[field] = clean(raw.get(f"__hyperlink__{header}") or raw.get(header, ""))
                else:
                    values[field] = clean(raw.get(header, "")) if header else ""
            record = LabelRecord(source_row=offset, **values)
            record.raw_fields = {header: clean(raw.get(header, "")) for header in self.headers}
            printable_values = [record.raw_fields.get(header, "") for header in self.headers if header != self.mapping_vars["link"].get()]
            is_group_header = record.name.upper().endswith("BOM") and not record.code and not record.supplier
            if any(printable_values) and not is_group_header:
                self.records.append(record)
        self.table_headers = [h for h in self.headers if h != self.mapping_vars["link"].get()]
        tree_columns = ["row"] + self.table_headers
        self.tree.configure(columns=tree_columns)
        self.tree.heading("row", text="Excel行")
        self.tree.column("row", width=70, minwidth=55, anchor="w")
        for header in self.table_headers:
            self.tree.heading(header, text=header)
            self.tree.column(header, width=190 if header != self.mapping_vars["name"].get() else 280, minwidth=80, anchor="w")
        self.tree.delete(*self.tree.get_children())
        for idx, record in enumerate(self.records):
            self.tree.insert("", "end", iid=str(idx), values=[record.source_row] + [record.raw_fields.get(h, "") for h in self.table_headers])
        if self.records:
            self.tree.selection_set("0")
            self.tree.focus("0")
        self.save_preferences()
        self.show_preview()

    def selected_records(self, all_records: bool) -> list[LabelRecord]:
        if all_records:
            return self.records[:]
        selected = [self.records[int(iid)] for iid in self.tree.selection()]
        if not selected and self.records:
            selected = [self.records[0]]
        return selected

    def copies_for(self, record: LabelRecord) -> int:
        if self.use_excel_qty.get() and record.qty:
            try:
                return max(1, min(999, int(float(record.qty))))
            except ValueError:
                pass
        return max(1, min(999, int(self.fixed_copies.get())))

    def show_preview(self):
        if not self.records:
            self.preview_canvas.delete("all")
            return
        try:
            records = self.selected_records(False)
            inbound = self.normalize_inbound_date(show_error=False) or self.inbound_date.get().strip()
            image = render_label(records[0], self.template_name.get(), self.display_columns_for(records[0]), inbound).convert("RGB")
            cw = max(100, self.preview_canvas.winfo_width() - 24)
            ch = max(100, self.preview_canvas.winfo_height() - 24)
            image.thumbnail((cw, ch), Image.Resampling.LANCZOS)
            self.preview_photo = ImageTk.PhotoImage(image)
            self.preview_canvas.delete("all")
            self.preview_canvas.create_image(self.preview_canvas.winfo_width() // 2, self.preview_canvas.winfo_height() // 2, image=self.preview_photo)
        except Exception as exc:
            self.status_text.set(f"预览失败：{exc}")

    def build_jobs(self, records: list[LabelRecord]) -> tuple[list[bytes], int]:
        inbound = self.normalize_inbound_date(show_error=False)
        if not inbound:
            raise ValueError("入库日期格式不正确，请使用 YYYY-MM-DD。")
        jobs: list[bytes] = []
        label_count = 0
        for record in records:
            copies = self.copies_for(record)
            payload = image_to_zpl(render_label(record, self.template_name.get(), self.display_columns_for(record), inbound))
            jobs.extend([payload] * copies)
            label_count += copies
        return jobs, label_count

    def display_columns_for(self, record: LabelRecord) -> list[tuple[str, str]]:
        selected: list[tuple[str, str]] = []
        seen: set[str] = set()
        for var in self.display_column_vars:
            header = var.get()
            if not header or header == "<不打印>" or header in seen:
                continue
            seen.add(header)
            selected.append((header, record.raw_fields.get(header, "")))
        return selected or [("零件", record.name), ("料号", record.code), ("供应商", record.supplier)]

    def export_zpl(self, all_records: bool):
        records = self.selected_records(all_records)
        if not records:
            messagebox.showinfo(APP_NAME, "没有可导出的记录。")
            return
        path = filedialog.asksaveasfilename(
            title="保存斑马 ZPL 打印文件",
            defaultextension=".zpl",
            filetypes=[("ZPL 打印文件", "*.zpl"), ("文本文件", "*.txt")],
            initialfile="labels.zpl",
        )
        if not path:
            return
        try:
            jobs, count = self.build_jobs(records)
            Path(path).write_bytes(b"".join(jobs))
            self.status_text.set(f"已导出 {count} 张标签：{path}")
            messagebox.showinfo(APP_NAME, f"已生成 {count} 张标签的 ZPL 文件。\n\n{path}")
        except Exception as exc:
            self.show_error("导出失败", exc)

    def print_records(self, all_records: bool):
        records = self.selected_records(all_records)
        printer = self.printer_name.get()
        if not records:
            messagebox.showinfo(APP_NAME, "没有可打印的记录。")
            return
        if not printer:
            messagebox.showwarning(APP_NAME, "尚未选择打印机。请先安装斑马打印机驱动并刷新打印机列表。")
            return
        try:
            jobs, count = self.build_jobs(records)
            target_text = "全部记录" if all_records else "选中记录"
            if not messagebox.askyesno(APP_NAME, f"将向打印机发送 {count} 张标签。\n\n范围：{target_text}\n打印机：{printer}\n\n确认打印吗？"):
                return
            send_raw(printer, jobs)
            self.printer_name.set(printer)
            self.save_preferences()
            self.status_text.set(f"已向 {printer} 发送 {count} 张标签")
            messagebox.showinfo(APP_NAME, f"打印任务已发送，共 {count} 张。\n请扫描首张标签核对内容。")
        except Exception as exc:
            self.show_error("打印失败", exc)

    def refresh_printers(self):
        try:
            printers = list_printers()
            self.printer_combo["values"] = printers
            if self.printer_name.get() not in printers:
                zebra = next((p for p in printers if "zebra" in p.lower() or "zd" in p.lower() or "zt" in p.lower()), "")
                self.printer_name.set(zebra)
            if not printers:
                self.status_text.set("未检测到 Windows 打印机；可先用“导出 ZPL”测试。")
        except Exception as exc:
            self.status_text.set(f"读取打印机失败：{exc}")

    def save_preferences(self):
        self.config_data.update({
            "template": self.template_name.get(),
            "printer": self.printer_name.get(),
            "sheet": self.sheet_name.get(),
            "mapping": {field: var.get() for field, var in self.mapping_vars.items()},
        })
        try:
            save_config(self.config_data)
        except Exception:
            pass

    def show_error(self, title: str, exc: Exception):
        details = "".join(traceback.format_exception_only(type(exc), exc)).strip()
        self.status_text.set(f"{title}：{details}")
        messagebox.showerror(APP_NAME, f"{title}\n\n{details}")


def main():
    app = LabelBridgeApp()
    app.mainloop()


if __name__ == "__main__":
    main()
