from __future__ import annotations

import csv
import functools
import io
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import barcode
import qrcode
from barcode.writer import ImageWriter
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

try:
    import win32print
except ImportError:
    win32print = None


FIELDS = [
    ("code", "物料编码/PN"),
    ("name", "物料名称"),
    ("supplier", "供应商"),
    ("rev", "版本/REV"),
    ("lot", "批次/LOT"),
    ("sn", "序列号/SN"),
    ("link", "飞书记录链接"),
    ("qty", "打印数量"),
    ("firmware", "固件版本"),
]

ALIASES = {
    "code": ["样机编号", "自动编号", "料号", "物料料号", "bom料号", "物料号", "货号", "物料编码", "零件编码", "物料编号", "零件号", "pn", "pn码", "partno", "partnumber", "编码", "code"],
    "name": ["零件", "物料名称", "零件名称", "名称", "品名", "name"],
    "supplier": ["供应商", "厂家", "厂商", "supplier", "vendor"],
    "rev": ["版本", "物料版本", "设计版本", "rev", "revision"],
    "lot": ["批次", "批次号", "lot", "lotnumber"],
    "sn": ["sn", "序列号", "整机sn", "产品sn", "serialnumber"],
    "link": ["记录分享链接", "分享链接", "记录链接", "飞书链接", "链接", "recordlink", "url", "二维码内容", "二维码"],
    "qty": ["打印数量", "标签数量", "数量", "qty", "copies", "打印份数"],
    "firmware": ["固件版本", "软件版本", "firmware", "fw"],
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return " / ".join(part for part in (clean(item) for item in value) if part)
    if isinstance(value, dict):
        for key in ("text", "name", "value"):
            if key in value:
                return clean(value[key])
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized(value: Any) -> str:
    return re.sub(r"[\s_\-/（）()]+", "", clean(value)).lower()


def auto_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {header: normalized(header) for header in headers}
    result: dict[str, str] = {}
    for field_name, aliases in ALIASES.items():
        exact = [h for h, normalized_header in normalized_headers.items() if normalized_header in aliases]
        fuzzy = [h for h, normalized_header in normalized_headers.items() if any(alias in normalized_header for alias in aliases)]
        result[field_name] = (exact or fuzzy or [""])[0]
    return result


@dataclass
class LabelRecord:
    source_row: int
    code: str = ""
    name: str = ""
    supplier: str = ""
    rev: str = ""
    lot: str = ""
    sn: str = ""
    link: str = ""
    qty: str = ""
    firmware: str = ""
    raw_fields: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class LabelSettings:
    width_mm: float = 40.0
    height_mm: float = 30.0
    dpi: int = 300
    safe_margin_mm: float = 2.2

    @property
    def dots_per_mm(self) -> float:
        return self.dpi / 25.4

    @property
    def pixel_size(self) -> tuple[int, int]:
        return (
            max(1, round(self.width_mm * self.dots_per_mm)),
            max(1, round(self.height_mm * self.dots_per_mm)),
        )


QR_SIZE_MIN_MM = 5.0
QR_SIZE_MAX_MM = 22.0
LOGO_WIDTH_MIN_MM = 6.0
LOGO_WIDTH_MAX_MM = 28.0
MAX_TEXT_LINES = 5
CONTENT_FIELD_SLOTS = 4
HAND_VALUES = {
    "左手": "L", "右手": "R", "左": "L", "右": "R",
    "L": "L", "R": "R", "l": "L", "r": "R",
}


@dataclass(frozen=True)
class LabelLayout:
    show_barcode: bool = True
    show_qr: bool = True
    niimbot_mode: str = ""
    offset_x_mm: float = 0.0
    offset_y_mm: float = 0.0
    text_scale_percent: int = 100
    text_spacing_mm: float = 0.2
    barcode_height_mm: float = 6.8
    barcode_width_percent: int = 100
    qr_size_mm: float = 9.5
    text_offset_x_mm: float = 0.0
    text_offset_y_mm: float = 0.0
    barcode_offset_x_mm: float = 0.0
    barcode_offset_y_mm: float = 0.0
    qr_offset_x_mm: float = 0.0
    qr_offset_y_mm: float = 0.0
    text_alignment: str = "center"
    line_scale_percent: tuple[int, ...] = (100, 100, 100, 100, 100)
    line_offset_x_mm: tuple[float, ...] = (0.0, 0.0, 0.0, 0.0, 0.0)
    line_offset_y_mm: tuple[float, ...] = (0.0, 0.0, 0.0, 0.0, 0.0)
    show_logo: bool = False
    logo_width_mm: float = 10.0
    logo_offset_x_mm: float = 0.0
    logo_offset_y_mm: float = 0.0
    element_abs_mm: tuple[tuple[str, float, float], ...] = ()


def _resource_roots() -> list[Path]:
    roots: list[Path] = []
    meipass = getattr(sys, "_MEIPASS", None)
    if meipass:
        roots.append(Path(meipass))
    here = Path(__file__).resolve().parent
    roots.extend([here, here / "fonts", here / "assets"])
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).resolve().parent
        roots.extend([exe_dir, exe_dir / "fonts", exe_dir / "assets"])
    return roots


def bundled_font_path() -> Path | None:
    for root in _resource_roots():
        for candidate in (root / "fonts" / "led_board-7.ttf", root / "led_board-7.ttf"):
            if candidate.exists():
                return candidate
    return None


def bundled_logo_path() -> Path | None:
    for root in _resource_roots():
        for candidate in (root / "assets" / "prime-logo.png", root / "prime-logo.png"):
            if candidate.exists():
                return candidate
    return None


@functools.lru_cache(maxsize=1)
def _logo_source() -> Image.Image | None:
    path = bundled_logo_path()
    if not path:
        return None
    try:
        return Image.open(path).convert("L")
    except Exception:
        return None


def make_logo(width_px: int) -> Image.Image | None:
    source = _logo_source()
    if source is None or width_px < 8:
        return None
    height_px = max(4, round(width_px * source.height / source.width))
    resized = source.resize((int(width_px), int(height_px)), Image.Resampling.LANCZOS)
    return resized.point(lambda p: 0 if p < 160 else 255)


def _font_supports(font, text: str) -> bool:
    for char in text:
        if char.isspace():
            continue
        try:
            box = font.getbbox(char)
        except Exception:
            return False
        if not box or box[3] - box[1] <= 0:
            return False
    return True


def get_font(size: int, bold: bool = False, text: str = ""):
    size = max(8, size)
    led = bundled_font_path()
    if led:
        try:
            font = ImageFont.truetype(str(led), size=size)
            if not text or _font_supports(font, text):
                return font
        except Exception:
            pass
    windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    # Gothic/black faces threshold cleanly on thermal printers. YaHei Bold is an
    # antialiased UI font; its gray halo becomes speckle after 1-bit conversion.
    names = ["simhei.ttf", "msyhbd.ttc"] if bold else ["msyh.ttc", "simhei.ttf"]
    names += ["msyh.ttc", "simsun.ttc", "arialbd.ttf" if bold else "arial.ttf"]
    seen: set[str] = set()
    for name in names:
        if name in seen:
            continue
        seen.add(name)
        path = windows_fonts / name
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size=size)
            except Exception:
                pass
    return ImageFont.load_default()


def to_thermal_bitmap(image: Image.Image) -> Image.Image:
    """Hard-threshold to 1-bit. Dithering around glyphs looks blurry on Zebra."""
    if image.mode != "L":
        image = image.convert("L")
    return image.convert("1", dither=Image.Dither.NONE)


def fit_text(draw: ImageDraw.ImageDraw, text: str, max_width: int, start_size: int, min_size: int, bold: bool = False):
    for size in range(max(start_size, min_size), min_size - 1, -1):
        font = get_font(size, bold, text)
        box = draw.textbbox((0, 0), text, font=font)
        if box[2] - box[0] <= max_width:
            return font
    return get_font(min_size, bold, text)


def explicit_break_parts(text: str) -> list[str]:
    """Split on real newlines and the two-character \\n typed in custom fields."""
    normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n").replace("\\n", "\n")
    return normalized.split("\n")


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    """Wrap a single paragraph to max_width without shrinking the font."""
    if text == "":
        return [""]
    lines: list[str] = []
    current = ""
    for char in text:
        trial = current + char
        box = draw.textbbox((0, 0), trial, font=font)
        if current and (box[2] - box[0]) > max_width:
            lines.append(current)
            current = "" if char.isspace() else char
        else:
            current = trial
    if current:
        lines.append(current)
    return lines or [""]


def field_display_rows(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    rows: list[str] = []
    for part in explicit_break_parts(text):
        rows.extend(wrap_text(draw, part, font, max_width))
    return rows or [text or ""]


def format_label_value(header: str, value: str) -> str:
    """Print field content only. 左右手 becomes L/R; no column titles on the label."""
    raw = clean(value)
    if header in {"左右手", "手性", "L/R", "LR"} or raw in HAND_VALUES:
        return HAND_VALUES.get(raw, raw[:1].upper() if raw else "")
    return raw


def _unified_text_scale(layout: LabelLayout) -> float:
    percent = layout.text_scale_percent or 100
    scales = layout.line_scale_percent or ()
    if percent == 100 and scales:
        percent = scales[0]
    return max(0.5, min(2.2, percent / 100))


def _line_scale(layout: LabelLayout, index: int) -> float:
    return _unified_text_scale(layout)


def _line_offset_mm(layout: LabelLayout, index: int) -> tuple[float, float]:
    xs = layout.line_offset_x_mm or ()
    ys = layout.line_offset_y_mm or ()
    x = xs[index] if index < len(xs) else 0.0
    y = ys[index] if index < len(ys) else 0.0
    if index == 0:
        x += layout.text_offset_x_mm
        y += layout.text_offset_y_mm
    return x, y


def make_qr(payload: str, target_size: int) -> Image.Image | None:
    if not clean(payload):
        return None
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=1, border=2)
    qr.add_data(clean(payload))
    qr.make(fit=True)
    raw = qr.make_image(fill_color="black", back_color="white").convert("1")
    size = max(raw.width, int(target_size))
    return raw.resize((size, size), Image.Resampling.NEAREST)


def make_code128(payload: str, width: int, height: int, dpi: int) -> Image.Image | None:
    payload = clean(payload)
    if not payload or width < 60 or height < 28:
        return None
    # Use integral printer dots for every module, including the quiet zones.
    # Fractional module widths followed by resizing can make long SNs unreadable.
    try:
        bits = barcode.get("code128", payload).build()[0]
    except Exception:
        return None
    module = width // (len(bits) + 20)
    if module < 1:
        raise ValueError("编号过长，当前条形码区域放不下；请加大标签宽度。")
    canvas = Image.new("1", (width, height), 1)
    draw = ImageDraw.Draw(canvas)
    start = (width - len(bits) * module) // 2
    for i, bit in enumerate(bits):
        if bit == "1":
            x = start + i * module
            draw.rectangle((x, 0, x + module - 1, height - 1), fill=0)
    return canvas


def _center_x(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> int:
    box = draw.textbbox((0, 0), text, font=font)
    return max(0, (width - (box[2] - box[0])) // 2)


def _render_label_with_regions(
    record: LabelRecord,
    display_columns: list[tuple[str, str]],
    settings: LabelSettings,
    printed_date: str = "",
    layout: LabelLayout | None = None,
    barcode_payload: str | None = None,
) -> tuple[Image.Image, dict[str, tuple[int, int, int, int]]]:
    layout = layout or LabelLayout()
    if layout.niimbot_mode:
        from compact_label import render_compact
        return render_compact(record, display_columns, settings, layout.niimbot_mode, barcode_payload)
    width, height = settings.pixel_size
    dpm = settings.dots_per_mm
    offset_x = round(layout.offset_x_mm * dpm)
    offset_y = round(layout.offset_y_mm * dpm)
    image = Image.new("L", (width, height), 255)
    draw = ImageDraw.Draw(image)
    if settings.width_mm < 20 or settings.height_mm < 15:
        # Small stock uses text only; the UI explicitly identifies this mode.
        margin = round(0.6 * dpm)
        date_text = clean(printed_date)
        field_budget = MAX_TEXT_LINES - (1 if date_text else 0)
        texts = [format_label_value(header, value) for header, value in display_columns[:field_budget] if clean(value)]
        if date_text:
            texts.append(date_text)
        texts = texts or [record.code or record.name or "未命名"]
        spacing = round(layout.text_spacing_mm * dpm)
        row_height = (height - 2 * margin - spacing * (len(texts) - 1)) // len(texts)
        if row_height < 9:
            raise ValueError("小标签内容过多：请减少打印字段或关闭日期。")
        y = margin
        regions: dict[str, tuple[int, int, int, int]] = {}
        for index, text in enumerate(texts):
            scale = _line_scale(layout, index)
            font = fit_text(draw, text, width - 2 * margin, min(round(24 * scale), row_height), 8, False)
            box = draw.textbbox((0, 0), text, font=font)
            tw, th = box[2] - box[0], box[3] - box[1]
            if tw > width - 2 * margin or th > row_height:
                raise ValueError("小标签文字放不下，请减少文字或打印字段。")
            extra_x, extra_y = _line_offset_mm(layout, index)
            x = margin if layout.text_alignment == "left" else width - margin - tw if layout.text_alignment == "right" else (width - tw) // 2
            x = max(margin, min(width - margin - tw, x + round((layout.offset_x_mm + extra_x) * dpm)))
            line_y = max(margin, min(height - margin - th, y + round(extra_y * dpm)))
            draw.text((x - box[0], line_y - box[1]), text, font=font, fill=0)
            regions[f"text_{index}"] = (x, line_y, x + tw, line_y + th)
            y += row_height + spacing
        if regions:
            first = regions["text_0"]
            regions["text"] = first
        return to_thermal_bitmap(image), regions
    # Do not stroke a frame around the bitmap. On black labels with white
    # ribbon that outline prints as a white box around the whole sticker.
    compact = settings.height_mm <= 20
    margin_mm = 1.6 if compact else settings.safe_margin_mm
    margin = max(round(margin_mm * dpm), 4 if compact else 6)
    content_width = width - margin * 2
    date_text = clean(printed_date)
    field_budget = MAX_TEXT_LINES - (1 if date_text else 0)
    lines = [(clean(header), format_label_value(header, value) or "—") for header, value in display_columns[:field_budget]]
    if date_text:
        lines.append(("", date_text))
    if not lines:
        lines = [("", record.name or record.code or "未命名")]

    qr_payload = record.link or record.code or record.sn
    logo_image = None
    if layout.show_logo:
        logo_width = round(max(LOGO_WIDTH_MIN_MM, min(LOGO_WIDTH_MAX_MM, layout.logo_width_mm)) * dpm)
        logo_width = max(8, min(logo_width, width - 2 * margin))
        logo_image = make_logo(logo_width)
    text_spacing = max(0, round(max(0.0, min(0.8, layout.text_spacing_mm)) * dpm))
    alignment = layout.text_alignment if layout.text_alignment in {"left", "center", "right"} else "center"
    graphic_reserve = 0
    if layout.show_qr and qr_payload:
        graphic_reserve = max(graphic_reserve, round(max(QR_SIZE_MIN_MM, min(QR_SIZE_MAX_MM, layout.qr_size_mm)) * dpm))
    if layout.show_barcode and (barcode_payload or record.code or record.sn):
        graphic_reserve = max(graphic_reserve, round(max(4.0, min(14.0, layout.barcode_height_mm)) * dpm))
    if graphic_reserve:
        graphic_reserve += max(2, round(0.3 * dpm))
    scale = _unified_text_scale(layout)
    font_size = max(8, round((2.6 if compact else 3.4) * dpm * scale))
    sample_font = get_font(font_size, False, "字")
    sample_box = draw.textbbox((0, 0), "字", font=sample_font)
    if sample_box[2] - sample_box[0] > content_width:
        sample_font = fit_text(draw, "字", content_width, font_size, 8, False)
        font_size = int(getattr(sample_font, "size", font_size) or font_size)

    prepared_fields: list[tuple[str, list[tuple[str, int, int]], object, int]] = []
    for header, value in lines:
        text = value
        font = get_font(font_size, False, text)
        rows = field_display_rows(draw, text, font, content_width)
        measured: list[tuple[str, int, int]] = []
        for row_text in rows:
            box = draw.textbbox((0, 0), row_text or " ", font=font)
            measured.append((row_text, box[2] - box[0] if row_text else 0, box[3] - box[1]))
        line_height = max((item[2] for item in measured), default=font_size)
        prepared_fields.append((header, measured, font, line_height))

    def is_hand_field(item: tuple) -> bool:
        header, rows, _font, _height = item
        first = rows[0][0] if rows else ""
        return first in {"L", "R"} or header in {"左右手", "手性", "L/R", "LR"}

    def aligned_x(block_width: int) -> int:
        if alignment == "left":
            return margin
        if alignment == "right":
            return width - margin - block_width
        return (width - block_width) // 2

    abs_pos = {name: (x_mm, y_mm) for name, x_mm, y_mm in layout.element_abs_mm}

    def placed_xy(name: str, default_x: int, default_y: int, box_w: int, box_h: int) -> tuple[int, int]:
        if name in abs_pos:
            x = round(abs_pos[name][0] * dpm)
            y = round(abs_pos[name][1] * dpm)
        else:
            x, y = default_x, default_y
        x = max(margin, min(width - margin - max(1, box_w), x))
        y = max(2, min(height - margin - max(1, box_h), y))
        return x, y

    def draw_row(text: str, font, x: int, y: int) -> None:
        box = draw.textbbox((0, 0), text or " ", font=font)
        draw.text((x - box[0], y - box[1]), text, font=font, fill=0)

    def place_field(index: int, rows: list[tuple[str, int, int]], font, start_x: int | None, start_y: int) -> tuple[int, int, int, int]:
        extra_x, extra_y = _line_offset_mm(layout, index)
        pieces: list[tuple[str, int, int, int, int]] = []
        y = start_y + round(extra_y * dpm)
        for row_text, tw, th in rows:
            x = (start_x if start_x is not None else aligned_x(max(tw, 1))) + offset_x + round(extra_x * dpm)
            pieces.append((row_text, x, y, tw, th))
            y += th + text_spacing
        origin_x, origin_y = pieces[0][1], pieces[0][2]
        pinned_x, pinned_y = placed_xy(
            f"text_{index}", origin_x, origin_y,
            max(item[3] for item in pieces),
            max(1, pieces[-1][2] + pieces[-1][4] - origin_y),
        )
        dx, dy = pinned_x - origin_x, pinned_y - origin_y
        xs: list[int] = []
        ys: list[int] = []
        bottoms: list[int] = []
        rights: list[int] = []
        for row_text, x, row_y, tw, th in pieces:
            x, row_y = x + dx, row_y + dy
            draw_row(row_text, font, x, row_y)
            xs.append(x)
            ys.append(row_y)
            rights.append(x + max(tw, 1))
            bottoms.append(row_y + th)
        bound = (min(xs), min(ys), max(rights), max(bottoms))
        regions[f"text_{index}"] = bound
        return bound

    cursor_y = margin + offset_y
    if logo_image:
        cursor_y += logo_image.height + max(2, round(0.25 * dpm))
    regions = {}
    last_text_bottom = margin
    index = 0
    while index < len(prepared_fields):
        header, rows, font, line_height = prepared_fields[index]
        companion = (
            index + 1 < len(prepared_fields)
            and is_hand_field(prepared_fields[index + 1])
            and len(rows) == 1
            and len(prepared_fields[index + 1][1]) == 1
            and f"text_{index}" not in abs_pos
            and f"text_{index + 1}" not in abs_pos
        )
        if companion:
            _header2, rows2, font2, height2 = prepared_fields[index + 1]
            text_width = rows[0][1]
            width2 = rows2[0][1]
            gap = max(6, round(0.8 * dpm))
            group_width = text_width + gap + width2
            group_height = max(line_height, height2)
            group_x = aligned_x(group_width)
            bound = place_field(index, rows, font, group_x, cursor_y + (group_height - line_height) // 2)
            bound2 = place_field(index + 1, rows2, font2, group_x + text_width + gap, cursor_y + (group_height - height2) // 2)
            last_text_bottom = max(last_text_bottom, bound[3], bound2[3])
            cursor_y = last_text_bottom + text_spacing
            index += 2
            continue
        bound = place_field(index, rows, font, None, cursor_y)
        last_text_bottom = max(last_text_bottom, bound[3])
        cursor_y = bound[3] + text_spacing
        index += 1
    if "text_0" in regions:
        regions["text"] = regions["text_0"]

    bottom = height - margin
    available_graphic_height = max(16, bottom - last_text_bottom - max(2, round(0.3 * dpm)))
    qr_size = round(max(QR_SIZE_MIN_MM, min(QR_SIZE_MAX_MM, layout.qr_size_mm)) * dpm)
    qr_size = max(16, min(qr_size, width - 2 * margin, height - 2 * margin, available_graphic_height))
    qr_image = make_qr(qr_payload, qr_size) if (qr_payload and layout.show_qr) else None
    gap = max(4, round(0.55 * dpm)) if qr_image else 0
    barcode_width = content_width - (qr_image.width + gap if qr_image else 0)
    # 100% intentionally uses 80% of the slot, leaving room for direct
    # manipulation to enlarge the barcode up to 125% without covering the QR.
    available_barcode_width = barcode_width
    barcode_width = max(60, round(barcode_width * max(0.4, min(1.0, layout.barcode_width_percent / 125))))
    payload = barcode_payload or record.code or record.sn
    if payload:
        try:
            minimum_width = len(barcode.get("code128", clean(payload)).build()[0]) + 20
            barcode_width = max(barcode_width, min(available_barcode_width, minimum_width))
        except Exception:
            pass
    barcode_height = min(round(max(4.0, min(14.0, layout.barcode_height_mm)) * dpm), height - 2 * margin, available_graphic_height)
    barcode_image = make_code128(payload, barcode_width, barcode_height, settings.dpi) if (layout.show_barcode and payload) else None
    group_width = (barcode_image.width if barcode_image else 0) + (gap if barcode_image and qr_image else 0) + (qr_image.width if qr_image else 0)
    group_height = max(barcode_image.height if barcode_image else 0, qr_image.height if qr_image else 0)
    gx = (width - group_width) // 2 + offset_x
    gx = max(margin, min(width - margin - group_width, gx))
    gy = min(bottom - group_height, last_text_bottom + max(2, round(0.3 * dpm)))
    gy = max(margin, gy)
    if barcode_image:
        default_x = gx + round(layout.barcode_offset_x_mm * dpm)
        default_y = gy + (group_height - barcode_image.height) // 2 + round(layout.barcode_offset_y_mm * dpm)
        barcode_x, barcode_y = placed_xy("barcode", default_x, default_y, barcode_image.width, barcode_image.height)
        image.paste(barcode_image.convert("L"), (barcode_x, barcode_y))
        regions["barcode"] = (barcode_x, barcode_y, barcode_x + barcode_image.width, barcode_y + barcode_image.height)
        gx += barcode_image.width + (gap if qr_image else 0)
    if qr_image:
        default_x = gx + round(layout.qr_offset_x_mm * dpm)
        default_y = gy + (group_height - qr_image.height) // 2 + round(layout.qr_offset_y_mm * dpm)
        qr_x, qr_y = placed_xy("qr", default_x, default_y, qr_image.width, qr_image.height)
        image.paste(qr_image.convert("L"), (qr_x, qr_y))
        regions["qr"] = (qr_x, qr_y, qr_x + qr_image.width, qr_y + qr_image.height)
    if logo_image:
        default_x = margin + offset_x + round(layout.logo_offset_x_mm * dpm)
        default_y = max(2, round(0.35 * dpm)) + offset_y + round(layout.logo_offset_y_mm * dpm)
        logo_x, logo_y = placed_xy("logo", default_x, default_y, logo_image.width, logo_image.height)
        image.paste(logo_image.convert("L"), (logo_x, logo_y))
        regions["logo"] = (logo_x, logo_y, logo_x + logo_image.width, logo_y + logo_image.height)
    return to_thermal_bitmap(image), regions


def render_label(
    record: LabelRecord,
    display_columns: list[tuple[str, str]],
    settings: LabelSettings,
    printed_date: str = "",
    layout: LabelLayout | None = None,
    barcode_payload: str | None = None,
) -> Image.Image:
    return _render_label_with_regions(record, display_columns, settings, printed_date, layout, barcode_payload)[0]


def render_label_with_regions(
    record: LabelRecord,
    display_columns: list[tuple[str, str]],
    settings: LabelSettings,
    printed_date: str = "",
    layout: LabelLayout | None = None,
    barcode_payload: str | None = None,
) -> tuple[Image.Image, dict[str, tuple[int, int, int, int]]]:
    """Render the printable bitmap and selectable element bounds for the editor."""
    return _render_label_with_regions(record, display_columns, settings, printed_date, layout, barcode_payload)


def image_to_zpl(
    image: Image.Image,
    print_method: str = "thermal_transfer",
    media_sensing: str = "gap",
) -> bytes:
    mono = to_thermal_bitmap(image) if image.mode != "1" else image
    width, height = mono.size
    row_bytes = (width + 7) // 8
    data = bytearray()
    pixels = mono.load()
    for y in range(height):
        for byte_x in range(row_bytes):
            value = 0
            for bit in range(8):
                x = byte_x * 8 + bit
                if x < width and pixels[x, y] == 0:
                    value |= 1 << (7 - bit)
            data.append(value)
    payload = data.hex().upper()
    method_command = "^MTT" if print_method == "thermal_transfer" else "^MTD"
    tracking_command = {"gap": "^MNY", "mark": "^MNM", "continuous": "^MNN"}.get(media_sensing, "^MNY")
    return f"^XA{method_command}{tracking_command}^PW{width}^LL{height}^LH0,0^FO0,0^GFA,{len(data)},{len(data)},{row_bytes},{payload}^FS^XZ\r\n".encode("ascii")


def zebra_calibration_zpl(print_method: str = "thermal_transfer", media_sensing: str = "gap") -> bytes:
    """Build a Zebra media/ribbon sensor calibration command."""
    method_command = "^MTT" if print_method == "thermal_transfer" else "^MTD"
    tracking_command = {"gap": "^MNY", "mark": "^MNM", "continuous": "^MNN"}.get(media_sensing, "^MNY")
    return f"^XA{method_command}{tracking_command}^XZ~JC\r\n".encode("ascii")


def _read_csv(path: Path) -> list[list[Any]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", newline="", encoding=encoding) as handle:
                return list(csv.reader(handle))
        except UnicodeDecodeError as exc:
            last_error = exc
    raise last_error or RuntimeError("无法识别 CSV 编码")


def load_local_rows(path_value: str, header_row: int = 1) -> tuple[list[str], list[dict[str, Any]], str]:
    path = Path(path_value)
    header_index = max(0, header_row - 1)
    suffix = path.suffix.lower()
    if suffix == ".json":
        package = json.loads(path.read_text(encoding="utf-8"))
        headers = [clean(header) for header in package.get("fields", [])]
        if "记录分享链接" not in headers:
            headers.append("记录分享链接")
        rows = []
        for item in package.get("records", []):
            values = {header: item.get("fields", {}).get(header, "") for header in headers}
            values["记录分享链接"] = item.get("record_share_link", "")
            rows.append(values)
        return headers, rows, path.stem
    if suffix == ".csv":
        matrix = _read_csv(path)
        if header_index >= len(matrix):
            raise ValueError("标题行超出 CSV 数据范围。")
        headers = [clean(value) or f"未命名列{index + 1}" for index, value in enumerate(matrix[header_index])]
        rows = [dict(zip(headers, row + [""] * (len(headers) - len(row)))) for row in matrix[header_index + 1:]]
        return headers, rows, path.stem
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    matrix = list(sheet.iter_rows())
    if header_index >= len(matrix):
        workbook.close()
        raise ValueError("标题行超出 Excel 数据范围。")
    headers = [clean(cell.value) or f"未命名列{index + 1}" for index, cell in enumerate(matrix[header_index])]
    rows: list[dict[str, Any]] = []
    for row in matrix[header_index + 1:]:
        values: dict[str, Any] = {}
        for index, header in enumerate(headers):
            cell = row[index] if index < len(row) else None
            values[header] = cell.value if cell else ""
            if cell and cell.hyperlink:
                values[f"__hyperlink__{header}"] = cell.hyperlink.target
        rows.append(values)
    workbook.close()
    return headers, rows, sheet.title


def records_from_rows(headers: list[str], rows: list[dict[str, Any]], mapping: dict[str, str], first_data_row: int = 2) -> list[LabelRecord]:
    records: list[LabelRecord] = []
    for offset, raw in enumerate(rows, start=first_data_row):
        values: dict[str, str] = {}
        for field_name, _label in FIELDS:
            header = mapping.get(field_name, "")
            if field_name == "link" and header:
                values[field_name] = clean(raw.get(f"__hyperlink__{header}") or raw.get(header, ""))
            else:
                values[field_name] = clean(raw.get(header, "")) if header else ""
        record = LabelRecord(source_row=offset, **values)
        record.raw_fields = {header: clean(raw.get(header, "")) for header in headers}
        printable = [record.raw_fields.get(header, "") for header in headers if header != mapping.get("link")]
        group_header = record.name.upper().endswith("BOM") and not record.code and not record.supplier
        if any(printable) and not group_header:
            records.append(record)
    return records


def list_printers() -> list[str]:
    if win32print is None:
        return []
    flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
    return sorted({printer[2] for printer in win32print.EnumPrinters(flags)}, key=str.lower)


def zebra_printers(printers: list[str] | None = None) -> list[str]:
    return [name for name in (printers or list_printers()) if any(token in name.lower() for token in ("zebra", "zdesigner", "zd", "zt", "gk888"))]


def print_queue_jobs(printer_name: str) -> list[dict[str, Any]]:
    if win32print is None:
        return []
    handle = win32print.OpenPrinter(printer_name)
    try:
        return list(win32print.EnumJobs(handle, 0, 999, 1))
    finally:
        win32print.ClosePrinter(handle)


def problematic_print_jobs(printer_name: str) -> list[dict[str, Any]]:
    problem_bits = 0
    for name in (
        "JOB_STATUS_PAUSED", "JOB_STATUS_ERROR", "JOB_STATUS_DELETING",
        "JOB_STATUS_OFFLINE", "JOB_STATUS_PAPEROUT", "JOB_STATUS_BLOCKED_DEVQ",
        "JOB_STATUS_USER_INTERVENTION",
    ):
        problem_bits |= int(getattr(win32print, name, 0)) if win32print is not None else 0
    return [
        job for job in print_queue_jobs(printer_name)
        if int(job.get("Size") or 0) == 0 or int(job.get("Status") or 0) & problem_bits
    ]


def cancel_print_jobs(printer_name: str) -> int:
    if win32print is None:
        raise RuntimeError("缺少 Windows 打印组件 pywin32。")
    handle = win32print.OpenPrinter(printer_name)
    cancelled = 0
    try:
        for job in win32print.EnumJobs(handle, 0, 999, 1):
            win32print.SetJob(handle, int(job["JobId"]), 0, None, win32print.JOB_CONTROL_DELETE)
            cancelled += 1
    finally:
        win32print.ClosePrinter(handle)
    return cancelled


def send_raw(printer_name: str, jobs: list[bytes], document_name: str = "Crelabel") -> int:
    if win32print is None:
        raise RuntimeError("缺少 Windows 打印组件 pywin32。")
    handle = win32print.OpenPrinter(printer_name)
    try:
        job_id = int(win32print.StartDocPrinter(handle, 1, (document_name, None, "RAW")))
        if not job_id:
            raise RuntimeError("Windows 未能创建打印任务。")
        try:
            win32print.StartPagePrinter(handle)
            for payload in jobs:
                written = int(win32print.WritePrinter(handle, payload))
                if written != len(payload):
                    raise RuntimeError(f"打印指令发送不完整：{written}/{len(payload)} 字节。")
            win32print.EndPagePrinter(handle)
        finally:
            win32print.EndDocPrinter(handle)
    finally:
        win32print.ClosePrinter(handle)
    return job_id
